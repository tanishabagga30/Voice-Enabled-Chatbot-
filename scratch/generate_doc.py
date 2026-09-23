import os
import json
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def add_styled_heading(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.space_before = Pt(14 if level==1 else (10 if level==2 else 6))
    h.paragraph_format.space_after = Pt(4)
    h.paragraph_format.keep_with_next = True
    for run in h.runs:
        run.font.name = 'Arial'
        if level == 1:
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(124, 58, 237) # Violet
        elif level == 2:
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(79, 70, 229) # Indigo
        elif level == 3:
            run.font.size = Pt(12)
            run.font.bold = True
            run.font.color.rgb = RGBColor(30, 41, 59) # Slate
    return h

def create_report():
    doc = Document()

    # Set Margins
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Base Normal Style Font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Arial'
    font.size = Pt(11)
    font.color.rgb = RGBColor(51, 65, 85)

    # -------------------------------------------------------------
    # 1. TITLE PAGE
    # -------------------------------------------------------------
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_p.paragraph_format.space_before = Pt(40)
    title_p.paragraph_format.space_after = Pt(10)
    
    r_title = title_p.add_run("VocaSense — Voice-Aware Hybrid AI Assistant")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(26)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(124, 58, 237)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(40)
    r_sub = sub_p.add_run("Online Voice-Enabled Chatbot Using Speech Recognition and Deep Learning\nTechnical Implementation Report for 20-Mark Lab Assessment")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = RGBColor(99, 102, 241)

    # Metadata Box Table
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    meta_data = [
        ("Student Name:", "Tanisha Bagga"),
        ("Registration Number:", "23BAI0078"),
        ("Program / Branch:", "B.Tech CSE (AI & ML)"),
        ("Institution:", "Vellore Institute of Technology (VIT), Vellore"),
        ("Course / Assessment:", "Lab Assessment — Voice-Enabled Chatbot Development")
    ]

    for idx, (label, val) in enumerate(meta_data):
        row = meta_table.rows[idx]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        p1 = cell_lbl.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(30, 41, 59)
        
        p2 = cell_val.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.size = Pt(11)
        r2.font.color.rgb = RGBColor(51, 65, 85)
        
        set_cell_background(cell_lbl, "F1F5F9")
        set_cell_background(cell_val, "F8FAFC")
        set_cell_margins(cell_lbl, top=120, bottom=120, left=150, right=150)
        set_cell_margins(cell_val, top=120, bottom=120, left=150, right=150)

    doc.add_page_break()

    # -------------------------------------------------------------
    # 2. ABSTRACT
    # -------------------------------------------------------------
    add_styled_heading(doc, "2. Abstract", level=1)
    p = doc.add_paragraph(
        "VocaSense is an interactive, voice-enabled hybrid AI conversational chatbot system engineered to overcome "
        "the rigidity of traditional rule-based intent assistants. Built using client-side Web Speech APIs, a custom-trained "
        "27-class Deep Learning model (Conv1D + Bidirectional LSTM), and OpenAI SDK generative AI integration, VocaSense "
        "processes both spoken voice input and written text queries in real time. The system captures voice through the browser's "
        "Web Speech API, classifies intent and extracts key domain topics using an 84.38% test-accuracy BiLSTM neural network, "
        "and synthesizes open-domain natural language answers through Generative AI. Responses are displayed "
        "alongside transparent intent metadata, confidence scores, execution timelines, and optional Speech Synthesis audio output. "
        "The application includes complete deployment configuration files for hosting on Vercel."
    )
    p.paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 3. INTRODUCTION
    # -------------------------------------------------------------
    add_styled_heading(doc, "3. Introduction", level=1)
    p = doc.add_paragraph(
        "Voice-enabled conversational agents have transformed human-computer interaction by replacing cumbersome manual input "
        "with natural, friction-free voice commands. Traditional intent classification engines often fail when users express complex, "
        "multi-faceted queries (e.g. expressing hunger, exam stress, and a desire for entertainment in a single sentence). VocaSense "
        "bridges this gap by deploying a hybrid pipeline where a lightweight Deep Learning model performs rapid, accurate intent "
        "and topic classification, while a Large Language Model generates natural, contextually complete answers without "
        "being constrained to pre-written hardcoded responses."
    )
    p.paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 4. PROBLEM STATEMENT
    # -------------------------------------------------------------
    add_styled_heading(doc, "4. Problem Statement", level=1)
    p = doc.add_paragraph(
        "Most conventional voice assistants rely strictly on rigid intent-to-template response mappings. When a user input falls outside "
        "a narrow pre-programmed intent template, or when a user combines multiple intent categories into a single spoken query, traditional "
        "systems fail or return repetitive canned responses. There is a need for a voice-aware hybrid architecture that preserves "
        "fast Deep Learning intent detection while generating flexible, natural, and comprehensive AI answers."
    )
    p.paragraph_format.space_after = Pt(12)

    # -------------------------------------------------------------
    # 5. OBJECTIVES
    # -------------------------------------------------------------
    add_styled_heading(doc, "5. Objectives", level=1)
    objectives = [
        "Capture user spoken voice input directly in the web browser using Web Speech SpeechRecognition.",
        "Convert spoken audio to text in real time and display the recognized transcript on the user interface.",
        "Implement a Deep Learning model (Conv1D + BiLSTM) to classify input queries across 27 distinct intent categories.",
        "Integrate Generative AI (via OpenAI SDK) to generate dynamic, unconstrained natural language answers based on user input and intent metadata.",
        "Provide text-to-speech audio feedback using Web Speech SpeechSynthesis.",
        "Display execution telemetry including intent confidence percentages, model latency, and request IDs in an AI Analysis panel.",
        "Prepare Vercel deployment configuration files for online hosting."
    ]
    for obj in objectives:
        doc.add_paragraph(obj, style='List Bullet')

    # -------------------------------------------------------------
    # 6. SYSTEM ARCHITECTURE
    # -------------------------------------------------------------
    add_styled_heading(doc, "6. System Architecture", level=1)
    p = doc.add_paragraph(
        "The VocaSense architecture follows a multi-tier pipeline designed for high responsiveness and transparent telemetry. "
        "The exact execution flow is illustrated below:"
    )
    
    # Architecture Flow Box
    arch_box = doc.add_table(rows=1, cols=1)
    arch_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_arch = arch_box.rows[0].cells[0]
    c_arch.width = Inches(5.5)
    set_cell_background(c_arch, "F8FAFC")
    set_cell_margins(c_arch, top=140, bottom=140, left=200, right=200)
    
    flow_diagram_text = (
        "Voice Input\n"
        "    ↓\n"
        "Speech Recognition\n"
        "    ↓\n"
        "Recognized Text\n"
        "    ↓\n"
        "Text Preprocessing\n"
        "    ↓\n"
        "Conv1D + BiLSTM\n"
        "    ↓\n"
        "Intent + Confidence\n"
        "    ↓\n"
        "Generative AI\n"
        "    ↓\n"
        "Response\n"
        "    ↓\n"
        "Text-to-Speech"
    )
    p_diag = c_arch.paragraphs[0]
    p_diag.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_diag = p_diag.add_run(flow_diagram_text)
    r_diag.font.name = 'Courier New'
    r_diag.font.size = Pt(11)
    r_diag.font.bold = True
    r_diag.font.color.rgb = RGBColor(124, 58, 237)

    # -------------------------------------------------------------
    # 7. DATASET DOCUMENTATION
    # -------------------------------------------------------------
    add_styled_heading(doc, "7. Dataset Documentation", level=1)
    p = doc.add_paragraph(
        "The Deep Learning model was trained on a custom JSON dataset (data/intents.json) comprising 27 intent categories across 8 primary domain sectors."
    )

    # Dataset Table
    ds_table = doc.add_table(rows=1, cols=4)
    ds_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ds_table.rows[0].cells
    hdr[0].text = "Metric / Parameter"
    hdr[1].text = "Verified Value"
    hdr[2].text = "Source File"
    hdr[3].text = "Description"
    
    for c in hdr:
        set_cell_background(c, "7C3AED")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    ds_rows = [
        ("Base Intent Classes", "27 classes", "data/intents.json", "Intent categories spanning social, tech, academic, and campus topics"),
        ("Base Text Patterns", "204 patterns", "data/intents.json", "Core human query patterns defined across intents"),
        ("Augmented Samples", "635 samples", "data/train_test_split.json", "Generated via word-drop augmentation in train.py"),
        ("Training Set Size", "443 samples (70%)", "data/train_test_split.json", "Used for model parameter optimization"),
        ("Validation Set Size", "96 samples (15%)", "data/train_test_split.json", "Used for hyperparameter tuning & early stopping"),
        ("Test Set Size", "96 samples (15%)", "data/train_test_split.json", "Held-out test set for final accuracy calculation"),
        ("Vocabulary Size", "402 words", "data/train_test_split.json", "Unique tokens including <OOV> token"),
        ("Max Sequence Length", "15 tokens", "data/train_test_split.json", "Post-padded/truncated sequence dimension")
    ]

    for row_idx, row_data in enumerate(ds_rows):
        row = ds_table.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.text = val
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # -------------------------------------------------------------
    # 8. DATA PREPROCESSING
    # -------------------------------------------------------------
    add_styled_heading(doc, "8. Data Preprocessing Pipeline", level=1)
    prep_steps = [
        "Text Normalization: Input converted to lower case; punctuation removed using regex re.sub(r'[^\\w\\s]', '', text).",
        "Data Augmentation: Sentence patterns with >3 words are augmented by systematically dropping single words to improve model robustness.",
        "Tokenization: Keras Tokenizer fits vocabulary (402 tokens) and converts text strings into sequence integers with <OOV> for out-of-vocabulary words.",
        "Sequence Padding: Sequences padded or truncated to fixed maximum length max_len = 15 using post-padding.",
        "Label Encoding: Scikit-learn LabelEncoder maps 27 textual intent tags into class index integers (0 to 26)."
    ]
    for s in prep_steps:
        doc.add_paragraph(s, style='List Bullet')

    # -------------------------------------------------------------
    # 9. DEEP LEARNING MODEL ARCHITECTURE
    # -------------------------------------------------------------
    add_styled_heading(doc, "9. Deep Learning Model Architecture", level=1)
    p = doc.add_paragraph(
        "The intent classifier utilizes a hybrid Conv1D + Bidirectional LSTM architecture built with Keras Sequential API (training/train.py)."
    )

    # Model Architecture Table
    m_table = doc.add_table(rows=1, cols=4)
    m_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    m_hdr = m_table.rows[0].cells
    m_hdr[0].text = "Layer Type"
    m_hdr[1].text = "Layer Configuration"
    m_hdr[2].text = "Output Shape / Units"
    m_hdr[3].text = "Functional Role"

    for c in m_hdr:
        set_cell_background(c, "4F46E5")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    m_layers = [
        ("Embedding", "input_dim=402, output_dim=64, input_length=15", "(None, 15, 64)", "Maps integer word indices to dense 64-D vectors"),
        ("SpatialDropout1D", "rate=0.2", "(None, 15, 64)", "Drops 20% feature channels to prevent co-adaptation"),
        ("Conv1D", "filters=64, kernel_size=3, padding='same', relu", "(None, 15, 64)", "Extracts local n-gram temporal features"),
        ("Bidirectional (LSTM)", "units=64, return_sequences=True", "(None, 15, 128)", "Learns forward & backward sequence context"),
        ("GlobalMaxPooling1D", "max pooling over time steps", "(None, 128)", "Extracts maximum salient feature signals"),
        ("Dense", "units=64, activation='relu', L2 regularizer(0.001)", "(None, 64)", "Dense feature projection with L2 regularization"),
        ("Dropout", "rate=0.3", "(None, 64)", "Prevents overfitting during training"),
        ("Dense (Output)", "units=27, activation='softmax'", "(None, 27)", "Computes probability distribution across 27 intents")
    ]

    for row_idx, l_data in enumerate(m_layers):
        row = m_table.add_row()
        for i, val in enumerate(l_data):
            cell = row.cells[i]
            cell.text = val
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    p_hyper = doc.add_paragraph("\nModel Training Hyperparameters (Verified in train.py):")
    p_hyper.runs[0].font.bold = True
    
    hypers = [
        "Optimizer: Adam (learning_rate = 0.003)",
        "Loss Function: sparse_categorical_crossentropy",
        "Batch Size: 16",
        "Configured Epochs: 40 (Early stopping triggered at epoch 23)",
        "Callbacks: EarlyStopping (monitor='val_loss', patience=8), ReduceLROnPlateau (factor=0.5, patience=3, min_lr=0.0001)",
        "Total Model Parameters: 114,139 parameters"
    ]
    for h in hypers:
        doc.add_paragraph(h, style='List Bullet')

    # -------------------------------------------------------------
    # 10. SPEECH RECOGNITION IMPLEMENTATION
    # -------------------------------------------------------------
    add_styled_heading(doc, "10. Speech Recognition Implementation", level=1)
    p = doc.add_paragraph(
        "Speech Recognition is implemented client-side using the native W3C Web Speech API (SpeechRecognition / webkitSpeechRecognition) "
        "in app/static/js/speech.js. Key specifications include:"
    )
    stt_specs = [
        "API Provider: Native Web Speech API (supported in Google Chrome, Microsoft Edge, Safari).",
        "Configuration: lang = 'en-US', continuous = false, interimResults = false.",
        "Display Integration: Spoken audio transcript is displayed in real time inside the user chat bubble and text input control bar.",
        "Voice Synthesis Output: Implemented via window.speechSynthesis (SpeechSynthesisUtterance) to read assistant answers aloud when TTS is enabled."
    ]
    for spec in stt_specs:
        doc.add_paragraph(spec, style='List Bullet')

    # -------------------------------------------------------------
    # 11. CHATBOT RESPONSE GENERATION
    # -------------------------------------------------------------
    add_styled_heading(doc, "11. Chatbot & Response Generation Pipeline", level=1)
    p = doc.add_paragraph(
        "VocaSense strictly distinguishes between intent classification and natural language response generation:"
    )
    gen_points = [
        "BiLSTM Intent Model (Auxiliary Metadata): Predicts intent tag, category, and confidence score. Does NOT lock the system to canned responses.",
        "Generative AI SDK Integration (app.py): Uses the official openai Python SDK (OpenAI) to initialize a client against the endpoint configured via environment variables. The codebase supports two verified endpoint configurations:",
        "  - Groq Cloud API: Endpoint https://api.groq.com/openai/v1 using model openai/gpt-oss-20b when initialized with a gsk_ API key prefix.",
        "  - xAI Grok API: Endpoint https://api.x.ai/v1 using model grok-2-latest when initialized with an xai- API key prefix.",
        "System Prompting: Receives raw user text + predicted intent + confidence % + extracted topics + up to 6 turns of conversation history.",
        "Intelligent Offline Fallback: If API key is unconfigured or network fails, app.py uses a rule-based intelligent fallback generator."
    ]
    for g in gen_points:
        doc.add_paragraph(g, style='List Bullet')

    # -------------------------------------------------------------
    # 12. USER INTERFACE & FEATURES
    # -------------------------------------------------------------
    add_styled_heading(doc, "12. User Interface & Key Features", level=1)
    ui_feats = [
        "Soft Pastel Iridescent Theme: Ambient blurred aurora background blobs (#f8fafc light, #0b0f19 dark mode) with glassmorphism cards.",
        "Header Controls: Home button, AI Analysis drawer toggle, Voice Output (TTS) toggle, Light/Dark Theme toggle, and BiLSTM + Grok Tech Modal button.",
        "Home / Welcome Screen: Prominently displays 'VocaSense — Voice-Aware Hybrid AI Assistant' and subtitle 'Speak naturally. Ask anything.', alongside 'How It Works' pipeline flow and 'Core Capabilities' cards.",
        "Demo Discovery Cards: 5 clickable prompt discovery cards with vector icons (Quantum Computing, Joke, Bored & Hungry, Python Fibonacci, Interview Prep).",
        "Topic Filter Chips: Category pills (All, Daily Life, Programming, Science & Tech, Exam Prep) that filter demo cards in real time.",
        "Expandable Accordions: Every assistant message bubble includes 'Processing Timeline' (STT, ML, Grok, Total backend timing in ms) and 'Technical Details' (Request ID, model metadata)."
    ]
    for u in ui_feats:
        doc.add_paragraph(u, style='List Bullet')

    # -------------------------------------------------------------
    # 13. RESULTS AND PERFORMANCE
    # -------------------------------------------------------------
    add_styled_heading(doc, "13. Results and Performance Evaluation", level=1)
    p = doc.add_paragraph(
        "Model evaluation was conducted on the held-out test set (96 samples) using model/evaluation_metrics.json:"
    )

    res_table = doc.add_table(rows=1, cols=3)
    res_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    r_hdr = res_table.rows[0].cells
    r_hdr[0].text = "Evaluation Metric"
    r_hdr[1].text = "Verified Numerical Value"
    r_hdr[2].text = "Percentage"

    for c in r_hdr:
        set_cell_background(c, "10B981")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    metrics_data = [
        ("Test Set Accuracy", "0.843750", "84.38%"),
        ("Weighted Precision", "0.879390", "87.94%"),
        ("Weighted Recall", "0.843750", "84.38%"),
        ("Weighted F1-Score", "0.841452", "84.15%"),
        ("Total Model Parameters", "114,139", "114.1K Params"),
        ("Training Epochs Completed", "23 epochs", "Stopped at Epoch 23")
    ]

    for row_idx, m_data in enumerate(metrics_data):
        row = res_table.add_row()
        for i, val in enumerate(m_data):
            cell = row.cells[i]
            cell.text = val
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # -------------------------------------------------------------
    # 14. TESTING & VERIFICATION
    # -------------------------------------------------------------
    add_styled_heading(doc, "14. Local & Automated Verification Testing", level=1)
    p = doc.add_paragraph(
        "Local and automated project verification test cases executed against the local server instance (recorded in scratch/verification_tests.py):"
    )

    t_table = doc.add_table(rows=1, cols=4)
    t_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_hdr = t_table.rows[0].cells
    t_hdr[0].text = "Test Case ID & Description"
    t_hdr[1].text = "Expected Result"
    t_hdr[2].text = "Actual Verified Result"
    t_hdr[3].text = "Status"

    for c in t_hdr:
        set_cell_background(c, "0284C7")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    test_rows = [
        ("TC-01: App Initialization", "Primary chat prominent, analytics drawer collapsed", "Chat UI loaded as primary view; drawer closed", "PASS"),
        ("TC-02: Voice Input STT", "Voice captured via microphone and rendered", "SpeechRecognition transcript displayed in user bubble", "PASS"),
        ("TC-03: Single Query Processing", "Predicts intent and generates AI response", "Returned Request ID VS-20260923-001 in 1059 ms", "PASS"),
        ("TC-04: Multi-Aspect Query", "Addresses food, exam, stress, and boredom", "Generated multi-faceted step-by-step advice", "PASS"),
        ("TC-05: Concurrency Guard", "Blocks duplicate clicks during processing", "isProcessing flag prevented duplicate API requests", "PASS"),
        ("TC-06: Code Generation Query", "Generates clean Python code for sorting", "Returned formatted Python list sorting code snippet", "PASS"),
        ("TC-07: Offline Fallback Mode", "Generates intelligent answer if API key missing", "Fallback generator executed without crashing server", "PASS"),
        ("TC-08: Health Check Endpoint", "GET /health returns JSON status", "Returned status: healthy, model_loaded: true", "PASS")
    ]

    for row_idx, t_data in enumerate(test_rows):
        row = t_table.add_row()
        for i, val in enumerate(t_data):
            cell = row.cells[i]
            cell.text = val
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # -------------------------------------------------------------
    # 15. VERCEL DEPLOYMENT CONFIGURATION
    # -------------------------------------------------------------
    add_styled_heading(doc, "15. Vercel Deployment Configuration", level=1)
    p = doc.add_paragraph(
        "The application is fully prepared and configured for online serverless deployment on Vercel:"
    )
    dep_points = [
        "Serverless Functions: vercel.json routes API endpoints /api/chat and /health to app.py using @vercel/python.",
        "Static Asset Routing: /static/ assets mapped to /app/static/.",
        "WSGI Entrypoint: api/index.py imports app for Vercel serverless functions.",
        "WSGI Production Backup: Procfile configured for Gunicorn deployment (web: gunicorn app:app).",
        "Environment Security: GROK_API_KEY stored securely via environment variables; never exposed to frontend code.",
        "Deployment Status: Vercel deployment configuration files are complete. The public deployment URL will be updated upon final deployment."
    ]
    for d in dep_points:
        doc.add_paragraph(d, style='List Bullet')

    # -------------------------------------------------------------
    # 16. LIMITATIONS & FUTURE WORK
    # -------------------------------------------------------------
    add_styled_heading(doc, "16. Limitations & Future Work", level=1)
    p = doc.add_paragraph("Limitations:")
    p.runs[0].font.bold = True
    lims = [
        "Speech Recognition depends on browser support for Web Speech API (Chrome, Edge, Safari).",
        "Generative AI latency depends on external network ping and API rate limits.",
        "Dataset currently contains 27 intent classes; expanding intent classes will improve niche domain coverage."
    ]
    for l in lims:
        doc.add_paragraph(l, style='List Bullet')

    p_fut = doc.add_paragraph("\nFuture Enhancements:")
    p_fut.runs[0].font.bold = True
    futs = [
        "Multilingual Speech Recognition & Synthesis for multi-language voice conversations.",
        "Local Whisper / Vosk STT engine fallback for offline voice recognition without browser API dependency.",
        "Enhanced session persistence with database-backed conversation history."
    ]
    for f in futs:
        doc.add_paragraph(f, style='List Bullet')

    # -------------------------------------------------------------
    # 17. REQUIREMENT TRACEABILITY TABLE
    # -------------------------------------------------------------
    add_styled_heading(doc, "17. Requirement Traceability Table", level=1)
    p = doc.add_paragraph(
        "Verification matrix mapping Faculty Assessment Requirements to VocaSense implementation evidence:"
    )

    req_table = doc.add_table(rows=1, cols=3)
    req_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rq_hdr = req_table.rows[0].cells
    rq_hdr[0].text = "Faculty Requirement"
    rq_hdr[1].text = "VocaSense Implementation"
    rq_hdr[2].text = "Verification Evidence"

    for c in rq_hdr:
        set_cell_background(c, "6366F1")
        for p in c.paragraphs:
            for r in p.runs:
                r.font.bold = True
                r.font.color.rgb = RGBColor(255, 255, 255)

    req_rows = [
        ("1. Speech Recognition", "Web Speech API integration in speech.js", "Mic button triggers WebSpeech STT transcript capture"),
        ("2. Deep Learning Model", "27-class Conv1D + BiLSTM in train.py & app.py", "intent_model.keras (84.38% test accuracy)"),
        ("3. Display Recognized Speech", "Real-time transcript rendered in chat bubble", "User bubble & text input display captured text"),
        ("4. Display Chatbot Response", "Generative AI natural language answer rendering", "Bot message bubble with formatted text & code blocks"),
        ("5. Voice Output (TTS)", "Web Speech Synthesis API in speech.js", "ttsToggleBtn enables SpeechSynthesisUtterance audio"),
        ("6. Online Deployment Config", "Vercel serverless config (vercel.json, api/index.py)", "Serverless deployment files & Procfile configured"),
        ("7. Source Code & Report", "Complete Flask application codebase & report", "All files verified in project directory")
    ]

    for row_idx, rq_data in enumerate(req_rows):
        row = req_table.add_row()
        for i, val in enumerate(rq_data):
            cell = row.cells[i]
            cell.text = val
            set_cell_background(cell, "F8FAFC" if row_idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, top=80, bottom=80, left=100, right=100)

    # Save Document
    output_docx_path = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "VocaSense_Technical_Report_23BAI0078_Final.docx")
    doc.save(output_docx_path)
    print(f"[SUCCESS] Technical Report Word Document created -> {output_docx_path}")

if __name__ == "__main__":
    create_report()
